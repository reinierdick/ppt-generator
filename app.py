from flask import Flask, request, send_file
from pptx import Presentation
import tempfile

app = Flask(__name__)

@app.route("/")
def home():
    return "PPT Generator Running"

@app.route("/create_ppt", methods=["POST"])
def create_ppt():

    data = request.json

    prs = Presentation()

    for slide_data in data["slides"]:

        slide = prs.slides.add_slide(
            prs.slide_layouts[1]
        )

        slide.shapes.title.text = slide_data["title"]

        slide.placeholders[1].text = "\n".join(
            slide_data.get("bullets", [])
        )

    tmp = tempfile.NamedTemporaryFile(
        delete=False,
        suffix=".pptx"
    )

    prs.save(tmp.name)

    return send_file(
        tmp.name,
        as_attachment=True,
        download_name="presentation.pptx",
        mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8080)
